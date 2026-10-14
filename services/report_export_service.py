from collections import defaultdict
import csv
from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO


def report_export_tables(report: dict) -> dict[str, list[dict]]:
    daily=defaultdict(lambda: {'Expected cashing (K)':Decimal('0')})
    for row in report['collections']:
        totals=daily[row['collection_date']]
        if row['expected_amount'] is None:
            totals['Expected cashing (K)']=None
        elif totals['Expected cashing (K)'] is not None:
            totals['Expected cashing (K)']+=row['expected_amount']
    # Collection headers include receipts whose bus allocation remains unresolved.
    receipts=defaultdict(lambda:Decimal('0'))
    for row in report['collection_headers']:
        receipts[row['collection_date']]+=row['total_amount']
    cashing=[{'Collection date':day,'Recorded cashing (K)':amount,
              'Expected cashing (K)':daily.get(day,{}).get('Expected cashing (K)')}
             for day,amount in sorted(receipts.items())]
    buses=[{'Bus':row['registration_number'],'Recorded cashing (K)':row['recorded_revenue'],
        'Confirmed allocation (K)':row['confirmed_revenue'],'Estimated allocation (K)':row['estimated_revenue'],
        'Expected cashing (K)':row['expected_cashing'],'Difference (K)':row['difference'],
        'Attributed costs (K)':row['attributed_costs'],'Contribution (K)':row['recorded_contribution'],
        'Active days':row['active_days'],'Deactivated days':row['deactivated_days'],
        'Unconfirmed days':row['unknown_days'],'Servicing events':row['servicing_events'],
        'Servicing costs (K)':row['servicing_costs'],'Current servicing due':row['next_servicing_due']}
        for row in report['buses']]
    driver=[]
    for row in report['salaries']:
        driver.append({'Driver':row['driver'],'Record':'Salary','Payroll month':row['payroll_month'],
            'Payment date':row['payment_date'],'Earned cost (K)':row['salary_cost'],
            'Net salary (K)':row['net_salary'],'Advance recovered (K)':row['advance_recovery']})
    for row in report['advances']:
        driver.append({'Driver':row['driver'],'Record':'Advance','Advance date':row['advance_date'],
            'Advance given (K)':row['amount'] if report['start']<=row['advance_date']<=report['end'] else Decimal('0'),
            'Outstanding at end (K)':row['outstanding']})
    owner=[]
    for row in report['loans']:
        owner.append({'Record':'Loan principal snapshot','Loan':row['loan_name'],'Lender':row['lender'],
            'Opening principal (K)':row['original_amount'],'Outstanding at end (K)':row['outstanding_principal']})
    for row in report['repayments']:
        owner.append({'Record':'Principal repayment','Date':row['repayment_date'],'Loan':row['loan_name'] or 'Unidentified',
            'Payment source':row['payment_source'],'Principal paid (K)':row['amount']})
    for row in report['owner']:
        owner.append({'Record':row['transaction_type'].replace('_',' ').title(),'Date':row['transaction_date'],
            'Owner amount (K)':row['amount'],'Description':row['description']})
    for kind,amount in report['balance']['other_payment_breakdown'].items():
        if kind=='OWNER_DRAWING':
            owner.append({'Record':'Owner drawings total','Business cash out (K)':amount})
    owner.append({'Record':'Owner liability at period end','Outstanding at end (K)':report['owner_balance']})
    return {'Daily Cashing Summary':cashing,'Bus Performance Summary':buses,
            'Expense Summary':report['categories'],'Owner and Loan Money Summary':owner,
            'Driver Salary and Advance Summary':driver}


def _safe_cell(value):
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')):
        return "'"+value
    return value


def csv_export(rows: list[dict]) -> bytes:
    headers=list(dict.fromkeys(key for row in rows for key in row))
    output=StringIO(newline='')
    writer=csv.DictWriter(output,fieldnames=headers)
    writer.writeheader()
    writer.writerows({key:_safe_cell(value) for key,value in row.items()} for row in rows)
    return output.getvalue().encode('utf-8-sig')


def excel_export(tables: dict[str,list[dict]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    workbook=Workbook()
    workbook.remove(workbook.active)
    for name,rows in tables.items():
        sheet=workbook.create_sheet(name[:31])
        headers=list(dict.fromkeys(key for row in rows for key in row))
        sheet.append(headers or ['No records for selected period'])
        for cell in sheet[1]:
            cell.font=Font(bold=True)
        for row in rows:
            sheet.append([_safe_cell(row.get(key)) for key in headers])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value,(Decimal,float,int)):
                    cell.number_format='#,##0.00;[Red]-#,##0.00'
                elif isinstance(cell.value,date):
                    cell.number_format='dd mmm yyyy'
        sheet.freeze_panes='A2'
        sheet.auto_filter.ref=sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width=min(38,max(18,len(str(column[0].value))+2))
    output=BytesIO()
    workbook.save(output)
    return output.getvalue()
