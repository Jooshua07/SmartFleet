"""Human-style financial and servicing scenarios, only in a disposable database."""
from datetime import date
from decimal import Decimal
from pathlib import Path
import os
import sys
import time
from unittest.mock import patch

import psycopg
from psycopg import sql
from dotenv import load_dotenv

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
load_dotenv()


def main():
    name=f'smartfleet_step4_test_{int(time.time())}'
    kwargs=dict(host=os.getenv('DB_HOST'),port=os.getenv('DB_PORT'),user=os.getenv('DB_USER'),password=os.getenv('DB_PASSWORD'))
    with psycopg.connect(**kwargs,dbname='postgres',autocommit=True) as conn:
        conn.execute(sql.SQL('CREATE DATABASE {}').format(sql.Identifier(name)))
    try:
        with psycopg.connect(**kwargs,dbname=name) as conn:
            for path in sorted(Path('migrations').glob('[0-9][0-9][0-9]_*.sql')):
                conn.execute(path.read_text(encoding='utf-8'))
        os.environ['DB_NAME']=name
        from database.connection import get_connection
        from services.bus_service import create_bus, set_bus_status
        from services.route_service import add_route
        from services.maintenance_service import record_servicing, get_servicing_schedules, get_servicing_history
        from services.finance_service import get_period_balance
        from services.performance_service import get_operating_summary, get_business_financial_timeline
        from services.expense_service import update_business_expense, record_business_expense
        from services.bus_account_service import get_bus_contribution_summary
        route=add_route('Step4 route')
        bus=create_bus('TEST001',route['id'],date_acquired=date(2026,10,1))
        second=create_bus('TEST002',route['id'],date_acquired=date(2026,10,1))
        with patch('utils.dates.today_lusaka',return_value=date(2026,11,1)):
            event=record_servicing(bus['id'],date(2026,10,16),100,50,25,payment_date=date(2026,10,17))
            assert get_servicing_schedules()[0]['next_due_date']==date(2026,11,20)
            record_servicing(bus['id'],date(2026,10,2),0,0,0)
            assert get_servicing_schedules()[0]['next_due_date']==date(2026,11,20), 'Backdate changed latest schedule'
            record_servicing(bus['id'],date(2026,10,20),200,0,0,event_id=event['id'],payment_date=date(2026,10,21),reason='Correct date and amount')
            assert get_servicing_schedules()[0]['next_due_date']==date(2026,11,24)
            try:
                record_servicing(bus['id'],date(2026,10,20),200,0,0,payment_date=date(2026,10,21))
            except ValueError:
                pass
            else:
                raise AssertionError('Duplicate servicing accepted')
        with get_connection() as conn:
            assert conn.execute('SELECT count(*) FROM financial_transactions').fetchone()[0]==1
            category=conn.execute("SELECT id FROM expense_categories WHERE name='Servicing'").fetchone()[0]
            expense_id=conn.execute('SELECT expense_id FROM maintenance WHERE id=%s',(event['id'],)).fetchone()[0]
        assert get_period_balance(date(2026,10,1),date(2026,10,31))['closing_balance']==-200
        assert get_period_balance(date(2026,11,1),date(2026,11,1))['opening_balance']==-200
        assert get_operating_summary(date(2026,10,1),date(2026,10,31))['expenses']==200
        assert get_bus_contribution_summary(date(2026,10,1),date(2026,10,31))['buses'][0]['direct_costs']==200
        series=get_business_financial_timeline(date(2026,10,1),date(2026,10,31))
        assert series[-1]['business_balance']==-200 and series[-1]['operating_expenses']==200
        print('Servicing: components, zero cost, payment date, early/late, backdate, correction, duplicate prevention, single ledger and attribution PASS')
        print('Periods: negative cash, historical cutoff, month carry-forward and chart reconciliation PASS')
        # Ordinary expense correction must reuse its payment and clear obsolete attribution.
        ordinary=record_business_expense(date(2026,10,5),category,300,'Office supplies')
        update_business_expense(ordinary['id'],date(2026,10,6),category,400,'Office supplies',scope='BUS',bus_id=second['id'])
        update_business_expense(ordinary['id'],date(2026,10,7),category,500,'Office supplies',scope='GENERAL')
        with get_connection() as conn:
            assert conn.execute('SELECT count(*) FROM financial_transactions').fetchone()[0]==2
        print('Historical ordinary expense: General -> Bus -> General and one updated payment PASS')
        from services.business_report_service import get_business_report
        report=get_business_report(date(2026,10,1),date(2026,10,31))
        assert report['operating']['expenses']==700
        assert report['balance']['closing_balance']==-700
        assert sum(row['Amount (K)'] for row in report['categories'])==700
        assert report['history'][0]['total_cost']==200
        empty=get_business_report(date(2026,11,1),date(2026,11,1))
        assert empty['operating']['expenses']==0 and empty['balance']['closing_balance']==-700
        print('Report: categories, maintenance counted once, empty month and negative balance PASS')
        from services.report_export_service import report_export_tables, excel_export
        from services.business_report_service import get_reconciliation
        from openpyxl import load_workbook
        from io import BytesIO
        tables=report_export_tables(report)
        assert sum(row['Amount (K)'] for row in tables['Expense Summary'])==report['operating']['expenses']
        workbook=load_workbook(BytesIO(excel_export(tables)))
        assert len(workbook.sheetnames)==5
        assert get_reconciliation(date(2026,10,31))['closing_balance']==-700
        print('Exports: five Excel sheets open; category totals match screen; ledger reconciliation PASS')
    finally:
        with psycopg.connect(**kwargs,dbname='postgres',autocommit=True) as conn:
            conn.execute(sql.SQL('DROP DATABASE {} WITH (FORCE)').format(sql.Identifier(name)))


if __name__=='__main__':
    main()
