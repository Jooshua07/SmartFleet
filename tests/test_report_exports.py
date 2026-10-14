import csv
from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO
import unittest

from openpyxl import load_workbook
from services.report_export_service import csv_export, excel_export


class ReportExportsTest(unittest.TestCase):
    def test_csv_and_excel_roundtrip(self):
        rows=[{'Bus':'AIF 9681','Amount (K)':Decimal('4000.25'),'Date':date(2026,10,1)},
              {'Bus':'AIH 2715','Amount (K)':Decimal('-800'),'Date':date(2026,10,4)}]
        exported=list(csv.DictReader(StringIO(csv_export(rows).decode('utf-8-sig'))))
        self.assertEqual(sum(Decimal(row['Amount (K)']) for row in exported),Decimal('3200.25'))
        book=load_workbook(BytesIO(excel_export({'Bus Performance Summary':rows})))
        sheet=book.active
        self.assertEqual(sheet['B2'].value,4000.25)
        self.assertEqual(sheet['B3'].value,-800)
        self.assertEqual(sheet['C2'].value.date(),date(2026,10,1))
        self.assertEqual(sheet.freeze_panes,'A2')

    def test_empty_and_formula_like_user_input(self):
        book=load_workbook(BytesIO(excel_export({'Empty':[]})))
        self.assertEqual(book.active['A1'].value,'No records for selected period')
        rows=[{'Description':'=1+1'}]
        book=load_workbook(BytesIO(excel_export({'Input':rows})))
        self.assertNotEqual(book.active['A2'].data_type,'f')
