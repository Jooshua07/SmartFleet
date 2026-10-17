import csv
import json
import os
import subprocess
import sys
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO
from uuid import uuid4

import psycopg
import pytest
from openpyxl import load_workbook

from database.connection import get_connection, shared_transaction
from services import maintenance_service
from smartfleet_web import operations, private_backups
from tests.integration.test_stage2_workflows import BASE, form, login, post
from tests.integration.test_stage3 import identifiers
from tests.integration.test_stage4 import reporting_data as reporting_data

pytestmark = pytest.mark.integration


def values(client, owner, path="/record-servicing", **overrides):
    return {
        **form(client, path),
        "bus_id": str(identifiers(owner)[2][0]),
        "completed_date": "2026-10-09",
        "labour": "200.25",
        "transport": "50.10",
        "other": "49.65",
        "payment_date": "2026-10-10",
        "notes": "Synthetic service",
        "reason": "Actual servicing",
        **overrides,
    }


def scalar(owner, query, params=()):
    with psycopg.connect(**owner) as connection:
        return connection.execute(query, params).fetchone()[0]


def test_paid_servicing_replay_correction_and_zero_cost(stage2_app, stage2_database):
    owner, _ = stage2_database
    client = login(stage2_app)
    data = values(client, owner)
    for _ in range(2):
        assert post(client, "/record-servicing", data).status_code == 303
    event = scalar(owner, "SELECT id FROM maintenance")
    assert scalar(owner, "SELECT total_cost FROM maintenance") == Decimal("300.00")
    assert scalar(owner, "SELECT amount FROM financial_transactions") == 300
    assert scalar(owner, "SELECT count(*) FROM expenses") == 1
    assert scalar(owner, "SELECT next_due_date FROM maintenance_schedules") == date(2026, 11, 13)
    assert scalar(owner, "SELECT actor FROM web_command_audit") == "browser-owner"
    with psycopg.connect(**owner) as connection:
        audit = connection.execute(
            "SELECT before_state,after_state FROM web_command_audit"
        ).fetchone()
        assert audit[0]["maintenance"] == []
        assert audit[1]["maintenance"][0]["total_cost"] == "300.00"
    fresh = client.get("/record-servicing", base_url=BASE).data.decode()
    assert 'name="labour" type="number" value="0"' in fresh
    path = f"/servicing/{event}/edit"
    changed = values(client, owner, path, labour="100", completed_date="2026-10-08")
    assert post(client, path, changed).status_code == 303
    assert scalar(owner, "SELECT count(*) FROM financial_transactions") == 1
    assert scalar(owner, "SELECT amount FROM financial_transactions") == Decimal("199.75")
    assert scalar(owner, "SELECT next_due_date FROM maintenance_schedules") == date(2026, 11, 12)
    zero = values(client, owner, path, labour="0", transport="0", other="0", payment_date="")
    assert post(client, path, zero).status_code == 303
    assert scalar(owner, "SELECT count(*) FROM financial_transactions") == 0
    assert scalar(owner, "SELECT expense_id FROM maintenance") is None
    paid_again = values(client, owner, path)
    assert post(client, path, paid_again).status_code == 303
    assert scalar(owner, "SELECT count(*) FROM expenses") == 1
    assert scalar(owner, "SELECT count(*) FROM financial_transactions") == 1


@pytest.mark.parametrize(
    "change",
    [
        {"bus_id": ""},
        {"bus_id": "99999999"},
        {"completed_date": "bad"},
        {"completed_date": "2099-01-01"},
        {"payment_date": ""},
        {"payment_date": "2099-01-01"},
        {"labour": "-1"},
        {"labour": "NaN"},
        {"labour": "Infinity"},
        {"labour": "1.001"},
        {"labour": "1000000000"},
        {"labour": ""},
        {"reason": ""},
        {"notes": "a" * 501},
    ],
)
def test_servicing_invalid_retains_fields_no_writes(stage2_app, stage2_database, change):
    owner, _ = stage2_database
    client = login(stage2_app)
    response = post(client, "/record-servicing", values(client, owner, **change))
    assert response.status_code == 400
    assert b"Synthetic service" in response.data or "notes" in change
    for table in (
        "maintenance",
        "expenses",
        "financial_transactions",
        "web_commands",
        "web_command_audit",
    ):
        assert scalar(owner, f"SELECT count(*) FROM {table}") == 0  # noqa: S608 -- fixed table whitelist.


def test_duplicate_date_stale_forms_and_schedule(stage2_app, stage2_database):
    owner, _ = stage2_database
    client = login(stage2_app)
    data = values(client, owner)
    stale = values(client, owner, completed_date="2026-10-07")
    assert post(client, "/record-servicing", data).status_code == 303
    assert post(client, "/record-servicing", stale).status_code == 409
    assert post(client, "/record-servicing", values(client, owner)).status_code == 400
    bus = identifiers(owner)[2][0]
    path = f"/servicing/schedule/{bus}"
    data = {
        **form(client, path),
        "next_due_date": "2026-12-01",
        "interval_days": "42",
        "reason": "Reviewed schedule",
    }
    assert post(client, path, data).status_code == 303
    assert post(client, path, data).status_code == 303
    # An earlier completion cannot overwrite the manually reviewed current due date.
    assert (
        post(
            client,
            "/record-servicing",
            values(
                client, owner, completed_date="2026-10-01", labour="0", other="0", transport="0"
            ),
        ).status_code
        == 303
    )
    assert scalar(owner, "SELECT next_due_date FROM maintenance_schedules") == date(2026, 12, 1)
    assert scalar(owner, "SELECT interval_days FROM maintenance_schedules") == 42
    assert (
        post(
            client,
            path,
            {
                **form(client, path),
                "next_due_date": "2026-12-01",
                "interval_days": "0",
                "reason": "Invalid",
            },
        ).status_code
        == 400
    )


def test_servicing_concurrent_retry_and_audit_rollback(stage2_app, stage2_database):
    owner, _ = stage2_database
    first, second = login(stage2_app), login(stage2_app)
    data = values(first, owner)
    other = {**data, "csrf_token": form(second, "/record-servicing")["csrf_token"]}
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(
            pool.map(
                lambda pair: post(pair[0], "/record-servicing", pair[1]).status_code,
                [(first, data), (second, other)],
            )
        )
    assert results == [303, 303]
    assert scalar(owner, "SELECT count(*) FROM maintenance") == 1
    assert scalar(owner, "SELECT count(*) FROM web_command_audit") == 1
    invalid = values(first, owner, completed_date="2026-10-07", reason="Audit failure\x00")
    assert post(first, "/record-servicing", invalid).status_code == 409
    assert scalar(owner, "SELECT count(*) FROM maintenance") == 1
    assert scalar(owner, "SELECT count(*) FROM financial_transactions") == 1


def test_servicing_deferred_commit_failure_rolls_back(stage2_app, stage2_database, monkeypatch):
    owner, _ = stage2_database
    original = maintenance_service.record_servicing

    def inconsistent(*args, **kwargs):
        result = original(*args, **kwargs)
        with get_connection() as connection:
            connection.execute(
                "UPDATE financial_transactions SET amount=amount+1 WHERE id=(SELECT financial_transaction_id FROM expenses WHERE id=(SELECT expense_id FROM maintenance WHERE id=%s))",
                (result["id"],),
            )
        return result

    monkeypatch.setattr(operations.maintenance, "record_servicing", inconsistent)
    client = login(stage2_app)
    assert post(client, "/record-servicing", values(client, owner)).status_code == 503
    for table in (
        "maintenance",
        "expenses",
        "financial_transactions",
        "web_commands",
        "web_command_audit",
    ):
        assert scalar(owner, f"SELECT count(*) FROM {table}") == 0  # noqa: S608


def test_servicing_restart_replay(stage2_app, stage2_database):
    owner, _ = stage2_database
    client = login(stage2_app)
    data = values(client, owner)
    assert post(client, "/record-servicing", data).status_code == 303
    code = """
import json,sys
from dotenv import dotenv_values
from werkzeug.security import generate_password_hash
from smartfleet_web import create_app
from tests.integration.test_stage2_workflows import login,form,post
e=dotenv_values('.env.flask.test.local')
assert e['SMARTFLEET_DB_NAME']=='smartfleet_flask_test'
db=dict(host=e['SMARTFLEET_DB_HOST'],port=e['SMARTFLEET_DB_PORT'],dbname=e['SMARTFLEET_DB_NAME'],user=e['SMARTFLEET_DB_USER'],password=e['SMARTFLEET_DB_PASSWORD'])
writer={**db,'user':e['SMARTFLEET_WRITE_DB_USER'],'password':e['SMARTFLEET_WRITE_DB_PASSWORD']}
app=create_app(dict(TESTING=True,SMARTFLEET_ENV='testing',SECRET_KEY=e['SMARTFLEET_SECRET_KEY'],OWNER_USERNAME='browser-owner',OWNER_PASSWORD_HASH=generate_password_hash('browser-synthetic-passphrase'),DB_SETTINGS=db,WRITE_DB_SETTINGS=writer,TRUSTED_HOSTS=['localhost']))
c=login(app)
v=json.loads(sys.stdin.read());v['csrf_token']=form(c,'/record-servicing')['csrf_token']
assert post(c,'/record-servicing',v).status_code==303
"""
    subprocess.run(  # noqa: S603 -- fixed interpreter and synthetic payload.
        [sys.executable, "-c", code],
        input=json.dumps(data),
        capture_output=True,
        text=True,
        timeout=90,
        check=True,
    )  # noqa: S603
    assert scalar(owner, "SELECT count(*) FROM maintenance") == 1


def test_reports_exports_reconcile_and_safe_cells(stage2_app, reporting_data, stage2_database):
    owner, _ = stage2_database
    client = login(stage2_app)
    with psycopg.connect(**owner) as connection:
        connection.execute(
            "UPDATE financial_transactions SET description='=SUM(1,2)' WHERE id=(SELECT min(id) FROM financial_transactions)"
        )
    args = "period=custom&custom=dates&start=2026-10-01&end=2026-10-31"
    response = client.get("/reports?" + args, base_url=BASE)
    assert response.status_code == 200, response.data.decode()[-1500:]
    assert b"3,600.00" in response.data and b"7,200.00" in response.data
    assert b"-3,600.00" in response.data
    response = client.get("/reports/export?" + args, base_url=BASE)
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data))
    assert len(workbook.sheetnames) == 10
    ledger = workbook["Ledger"]
    assert any(
        cell.value == "'=SUM(1,2)" and cell.data_type != "f" for row in ledger for cell in row
    )
    summary = list(workbook["Business Summary"].values)
    totals = dict(zip(summary[0], summary[1], strict=True))
    assert totals["Cash money_in"] == 3600
    assert totals["Cash money_out"] == 7200
    assert totals["Cash closing_balance"] == -3600
    csv_response = client.get("/reports/export?" + args + "&format=csv&table=Ledger", base_url=BASE)
    rows = list(csv.DictReader(StringIO(csv_response.data.decode("utf-8-sig"))))
    assert any(r["description"] == "'=SUM(1,2)" for r in rows)
    assert not any(r["transaction_date"].startswith("2026-11") for r in rows)
    assert (
        client.get("/reports/export?format=csv&table=../../.env", base_url=BASE).status_code == 400
    )
    assert (
        client.get(
            "/reports?period=custom&custom=dates&start=bad&end=2026-10-01", base_url=BASE
        ).status_code
        == 400
    )


def test_reminders_help_and_empty_reports(stage2_app, stage2_database, tmp_path):
    owner, _ = stage2_database
    stage2_app.config["PRIVATE_BACKUP_ROOT"] = tmp_path
    client = login(stage2_app)
    with shared_transaction(owner):
        bus = identifiers(owner)[2][0]
        maintenance_service.set_servicing_schedule(bus, date(2026, 10, 16))
    response = client.get("/reminders", base_url=BASE)
    assert response.status_code == 200 and b"Overdue" in response.data
    for path in ("/reports", "/servicing", "/backups", "/more", "/help?q=servicing"):
        assert client.get(path, base_url=BASE).status_code == 200
    assert (
        b"How do I record completed servicing?"
        in client.get("/help?q=completed+servicing", base_url=BASE).data
    )
    assert b"No matching topics" in client.get("/help?q=nothingmatches123", base_url=BASE).data


def test_private_backup_replay_download_and_real_isolated_restore(
    stage2_app, stage2_database, tmp_path
):
    owner, _ = stage2_database
    settings = stage2_app.config["DB_SETTINGS"]
    stage2_app.config["PRIVATE_BACKUP_ROOT"] = tmp_path
    client = login(stage2_app)
    data = form(client, "/backups")
    # Actual restore rehearsals are restricted to this named synthetic database.
    assert (
        owner["dbname"] == "smartfleet_flask_test"
        and owner["user"] == "smartfleet_flask_test_owner"
    )
    assert owner["host"] in {"localhost", "127.0.0.1", "::1"}
    # Restore meaningful financial history, not merely an empty schema.
    assert post(client, "/record-servicing", values(client, owner)).status_code == 303
    from tests.integration.test_complete_cashing import payload

    assert post(client, "/record-cashing", payload(client)).status_code == 303

    def database_fingerprint(connection):
        from psycopg import sql

        tables = connection.execute(
            "SELECT tablename FROM pg_tables WHERE schemaname='public' ORDER BY tablename"
        ).fetchall()
        return {
            table: connection.execute(
                sql.SQL(
                    "SELECT count(*),md5(string_agg(row_to_json(t)::text,'' ORDER BY row_to_json(t)::text)) FROM {} t"
                ).format(sql.Identifier(table))
            ).fetchone()
            for (table,) in tables
        }

    with psycopg.connect(**owner) as connection:
        full_before = database_fingerprint(connection)
        before = connection.execute(
            "SELECT md5(string_agg(row_to_json(b)::text,'' ORDER BY id)) FROM buses b"
        ).fetchone()[0]
        triggers = connection.execute(
            "SELECT count(*) FROM pg_trigger WHERE NOT tgisinternal"
        ).fetchone()[0]
    assert post(client, "/backups", data).status_code == 303
    assert post(client, "/backups", data).status_code == 303
    # A new app/session still uses the same on-disk completed request identity.
    from smartfleet_web import create_app

    restarted = create_app({**stage2_app.config, "PRIVATE_BACKUP_ROOT": tmp_path})
    retry = login(restarted)
    assert (
        post(
            retry, "/backups", {**data, "csrf_token": form(retry, "/backups")["csrf_token"]}
        ).status_code
        == 303
    )
    records = private_backups.list_archives(settings, tmp_path)
    assert len(records) == 1 and records[0]["actor"] == "browser-owner"
    key = records[0]["key"]
    response = client.get(f"/backups/{key}/download", base_url=BASE)
    assert response.status_code == 200 and response.data[:5] == b"PGDMP"
    response.close()
    archive, _ = private_backups.archive(settings, key, tmp_path)
    with psycopg.connect(**owner) as connection:
        assert database_fingerprint(connection) == full_before
        connection.execute("UPDATE buses SET notes='Synthetic post-backup change'")
    restore = private_backups._find_tool("pg_restore", private_backups.PG_RESTORE_CANDIDATES)
    env = {**os.environ, "PGPASSWORD": owner["password"], "PGCONNECT_TIMEOUT": "5"}
    env.pop("PGOPTIONS", None)
    result = subprocess.run(  # noqa: S603 -- fixed tool and guarded synthetic database.
        [
            restore,
            "--host",
            owner["host"],
            "--port",
            str(owner["port"]),
            "--username",
            owner["user"],
            "--dbname",
            owner["dbname"],
            "--clean",
            "--if-exists",
            "--single-transaction",
            "--exit-on-error",
            "--no-password",
            str(archive),
        ],
        env=env,
        capture_output=True,
        timeout=90,
    )  # noqa: S603 -- fixed tool, explicitly isolated test database.
    assert result.returncode == 0, result.stderr.decode()[:1500]
    with psycopg.connect(**owner) as connection:
        assert database_fingerprint(connection) == full_before
        assert (
            connection.execute(
                "SELECT md5(string_agg(row_to_json(b)::text,'' ORDER BY id)) FROM buses b"
            ).fetchone()[0]
            == before
        )
        assert (
            connection.execute("SELECT count(*) FROM pg_trigger WHERE NOT tgisinternal").fetchone()[
                0
            ]
            == triggers
        )
        assert (
            connection.execute(
                "SELECT count(*) FROM information_schema.tables WHERE table_schema='public' AND table_type='BASE TABLE'"
            ).fetchone()[0]
            == 41
        )
    # Tampered files and traversal references cannot be downloaded as trusted backups.
    with archive.open("ab") as file:
        file.write(b"corrupt")
    assert client.get(f"/backups/{key}/download", base_url=BASE).status_code == 404
    assert client.get("/backups/not-a-uuid/download", base_url=BASE).status_code == 404


def test_backup_failure_and_serialized_lock(stage2_app, stage2_database, tmp_path, monkeypatch):
    owner, _ = stage2_database
    settings = stage2_app.config["DB_SETTINGS"]

    def failed(*args, **kwargs):
        raise ValueError("Synthetic dump failure")

    monkeypatch.setattr(private_backups, "bounded_run", failed)
    with pytest.raises(ValueError, match="Synthetic"):
        private_backups.create_archive(settings, str(uuid4()), "browser-owner", tmp_path)
    assert private_backups.list_archives(settings, tmp_path) == []
    assert not list(tmp_path.rglob("*.partial"))
    with psycopg.connect(**owner, autocommit=True) as connection:
        connection.execute("SELECT pg_advisory_lock(hashtext('smartfleet-private-backup'))")
        with pytest.raises(ValueError, match="Another backup"):
            private_backups.create_archive(settings, str(uuid4()), "browser-owner", tmp_path)
    monkeypatch.setattr(
        private_backups.shutil, "disk_usage", lambda path: type("Usage", (), {"free": 0})()
    )
    with pytest.raises(ValueError, match="storage"):
        private_backups.create_archive(settings, str(uuid4()), "browser-owner", tmp_path)
